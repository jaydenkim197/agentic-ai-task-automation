from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT = PROJECT_ROOT / "docs" / "final-report" / "Final Report - Agentic AI-Based Task Automation System.docx"
DIAGRAM = Path(__file__).resolve().parent / "assets" / "system_architecture.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "EAF4EA"
WHITE = "FFFFFF"
BLACK = "000000"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic
    return run


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 15, 7),
        ("Heading 2", 13, BLUE, 11, 5),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.45)
        style.paragraph_format.first_line_indent = Inches(-0.2)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def add_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("AGENTIC AI TASK AUTOMATION  |  FINAL REPORT"), 8.5, True, MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    set_run(p.add_run("Final Report  •  Page "), 8.5, False, MUTED)
    add_page_field(p)


def add_title_page(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(15)
    set_run(kicker.add_run("FINAL PROJECT REPORT"), 11, True, BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run(
        title.add_run("Design and Implementation of an\nAgentic AI-Based Personal Task\nAutomation System"),
        26,
        True,
        INK,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_run(
        subtitle.add_run(
            "A lightweight Discord-controlled automation agent inspired by OpenClaw"
        ),
        13,
        False,
        MUTED,
        True,
    )

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    data = [
        ("Planning date", "February 19, 2026"),
        ("Implementation period", "June 20–21, 2026"),
        ("Development team", "One developer"),
        ("Target platform", "MacBook Air M1, macOS"),
        ("Project status", "MVP implemented and operational"),
    ]
    for row, (label, value) in zip(table.rows, data):
        set_cell_shading(row.cells[0], PALE_BLUE)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), 10, True, DARK_BLUE)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(value), 10, False, BLACK)
    set_table_geometry(table, [2500, 5300], 0)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    set_run(p.add_run("Prepared for final project submission"), 10, False, MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("June 21, 2026"), 10, True, MUTED)
    doc.add_page_break()


def add_callout(doc, title: str, text: str, fill=PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(title), 10.5, True, DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(text), 10.5, False, BLACK)
    set_table_geometry(table, [9360], 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.keep_together = True
        p.add_run(item)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.keep_together = True
        p.add_run(item)


def add_status_table(doc) -> None:
    rows = [
        ("Discord interface", "Complete", "Owner-only natural-language commands; status messages and chunked replies"),
        ("OCR automation", "Complete", "Korean/English text, LaTeX math, Markdown tables, figure descriptions"),
        ("Gmail management", "Complete", "Explicit date range, classification, safe Trash action, summary report"),
        ("Reply workflow", "Complete", "Draft ID with explicit send/cancel confirmation"),
        ("Daily briefing", "Complete", "Gmail, Sheets, and Calendar summarized directly in Discord"),
        ("24/7 operation", "Complete", "macOS launchd RunAtLoad + KeepAlive; AC sleep disabled"),
        ("Automated tests", "21 passed", "Security, routing, ranges, reports, OCR and controller behavior"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Component", "Status", "Evidence")):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 9.5, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for component, status, evidence in rows:
        cells = table.add_row().cells
        for i, text in enumerate((component, status, evidence)):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(text), 9.2, i == 1, DARK_BLUE if i == 1 else BLACK)
        if status in {"Complete", "21 passed"}:
            set_cell_shading(cells[1], PALE_GREEN)
    set_table_geometry(table, [2300, 1500, 5560])


def add_requirements_table(doc) -> None:
    rows = [
        ("Interface", "Discord natural-language messages", "Implemented"),
        ("File source", "OneDrive-mounted local directory only", "Implemented"),
        ("OCR output", "UTF-8 TXT beside source, `_ag` suffix", "Implemented"),
        ("OCR content", "Korean, English, math, tables, figures", "Implemented"),
        ("Gmail", "Date-bounded review, summary, safe Trash", "Implemented"),
        ("Email send", "Only after explicit user confirmation", "Implemented"),
        ("Calendar", "Read only", "Implemented"),
        ("Task records", "Google Sheets read only", "Implemented"),
        ("Memory", "No database or long-term conversation memory", "Implemented"),
        ("Security", "Owner restriction and prompt-injection defense", "Implemented"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Area", "Requirement", "Result")):
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 9.5, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for area, requirement, result in rows:
        cells = table.add_row().cells
        for i, text in enumerate((area, requirement, result)):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(text), 9.2, i == 2, DARK_BLUE if i == 2 else BLACK)
    set_table_geometry(table, [1800, 5760, 1800])


def add_test_table(doc) -> None:
    rows = [
        ("OCR: timetable PDF", "1 page", "TXT generated; UTF-8/page marker verified", "Pass"),
        ("OCR: engineering math HW#1", "1 page", "Korean and LaTeX formulas preserved", "Pass"),
        ("Overwrite prevention", "Repeated output path", "Existing `_ag.txt` blocked before write", "Pass"),
        ("Discord access control", "Authorized account", "Command received and processed", "Pass"),
        ("Gmail organization", "2026-06-11 to 2026-06-12", "10 reviewed; 7 promotional messages moved to Trash", "Pass"),
        ("Email report", "Gmail result", "`important_emails.md` generated", "Pass"),
        ("Reply approval", "Date-bounded email search", "Draft generated; cancel flow confirmed", "Pass"),
        ("Daily briefing", "Gmail + Sheets + Calendar", "Direct Discord response", "Pass"),
        ("Automated test suite", "21 tests", "All tests passed", "Pass"),
        ("Persistent operation", "launchd service", "Running with KeepAlive and RunAtLoad", "Pass"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Test", "Input / Scope", "Observed result", "Status")):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 9, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for test, scope, result, status in rows:
        cells = table.add_row().cells
        for i, text in enumerate((test, scope, result, status)):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 3 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(text), 8.7, i == 3, DARK_BLUE if i == 3 else BLACK)
        set_cell_shading(cells[3], PALE_GREEN)
    set_table_geometry(table, [2150, 2050, 3960, 1200])


def add_comparison_table(doc) -> None:
    rows = [
        ("OpenClaw", "Self-hosted personal agent", "Chat apps; broad tools; persistent context", "High flexibility; broader privilege and setup surface"),
        ("Zapier Agents", "Managed SaaS", "Specialized agents across 9,000+ apps", "Fast integration breadth; platform-dependent"),
        ("n8n AI", "Cloud or self-hosted workflows", "Visual orchestration and AI agent nodes", "Strong audit/human checkpoints; more workflow configuration"),
        ("Microsoft 365 Copilot", "Microsoft-managed service", "Work-context chat and agents across Microsoft 365", "Strong governance; centered on Microsoft ecosystem"),
        ("This project", "Local Python service on macOS", "Discord + OCR + Gmail + Sheets + Calendar", "Narrow scope with explicit permissions and deterministic guardrails"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("System", "Deployment", "Primary strength", "Project-relevant trade-off")):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), 8.8, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(text), 8.5, i == 0, DARK_BLUE if i == 0 else BLACK)
        if values[0] == "This project":
            for cell in cells:
                set_cell_shading(cell, PALE_GREEN)
    set_table_geometry(table, [1500, 1900, 2850, 3110])


def make_architecture_diagram() -> None:
    DIAGRAM.parent.mkdir(exist_ok=True)
    image = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(font_path, 28)
    bold = ImageFont.truetype(bold_path, 32)
    small = ImageFont.truetype(font_path, 23)

    boxes = [
        ((520, 45, 980, 135), "Discord Owner Interface", PALE_BLUE),
        ((520, 190, 980, 280), "Agent Controller", "DCE6F1"),
        ((520, 335, 980, 425), "GPT API\nInterpretation / Classification", "EDE7F6"),
        ((70, 510, 390, 625), "OCR Tool\nOneDrive → _ag.txt", "EAF4EA"),
        ((420, 510, 740, 625), "Gmail Tool\nReview / Trash / Draft", "FFF3CD"),
        ((770, 510, 1090, 625), "Sheets Tool\nRead-only tasks", "EAF4EA"),
        ((1120, 510, 1430, 625), "Calendar Tool\nRead-only events", "EAF4EA"),
        ((520, 720, 980, 825), "Results\nDiscord response / local report", PALE_BLUE),
    ]
    for (x1, y1, x2, y2), label, color in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=f"#{color}", outline=f"#{DARK_BLUE}", width=3)
        lines = label.split("\n")
        heights = [draw.textbbox((0, 0), line, font=bold if len(lines) == 1 else font)[3] for line in lines]
        total = sum(heights) + (len(lines) - 1) * 8
        y = (y1 + y2 - total) / 2
        for line in lines:
            use_font = bold if len(lines) == 1 else font
            bbox = draw.textbbox((0, 0), line, font=use_font)
            draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, fill=f"#{INK}", font=use_font)
            y += bbox[3] - bbox[1] + 8

    arrows = [
        ((750, 135), (750, 190)),
        ((750, 280), (750, 335)),
        ((750, 425), (230, 510)),
        ((750, 425), (580, 510)),
        ((750, 425), (930, 510)),
        ((750, 425), (1275, 510)),
        ((230, 625), (650, 720)),
        ((580, 625), (700, 720)),
        ((930, 625), (800, 720)),
        ((1275, 625), (850, 720)),
    ]
    for start, end in arrows:
        draw.line((*start, *end), fill=f"#{DARK_BLUE}", width=4)
        x, y = end
        draw.polygon([(x, y), (x - 10, y - 18), (x + 10, y - 18)], fill=f"#{DARK_BLUE}")
    draw.text((45, 850), "Model output is validated before deterministic tools execute.", fill=f"#{MUTED}", font=small)
    image.save(DIAGRAM)


def add_source(doc, number: int, title: str, url: str, note: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(f"[{number}] {title}. "), 9.3, True, BLACK)
    set_run(p.add_run(note + " "), 9.3, False, BLACK)
    run = p.add_run(url)
    set_run(run, 9.0, False, BLUE)
    run.underline = True


def build_report() -> None:
    make_architecture_diagram()
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc.sections[0])
    add_title_page(doc)

    doc.add_heading("Executive Summary", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "This project designed and implemented a lightweight personal AI agent that converts natural-language Discord messages into controlled automation actions. "
        "The system was built from scratch in Python rather than installing OpenClaw. It runs continuously on a MacBook Air M1 and integrates a hosted GPT model with a restricted set of tools for PDF OCR, Gmail management, Google Sheets task reading, Google Calendar inspection, and daily productivity briefings."
    )
    add_callout(
        doc,
        "Final result",
        "The MVP is operational. The complete automated test suite reports 21 passed tests, real OCR files were generated successfully, Gmail organization processed a specified date range, and the Discord bot remains online through a macOS launchd service.",
        PALE_GREEN,
    )
    add_status_table(doc)

    doc.add_heading("1. Project Background and Motivation", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The project was planned on February 19, 2026 after the developer encountered early-2026 demonstrations of OpenClaw and overseas creators using always-on AI agents to automate routine digital work. These demonstrations suggested a transition from passive question-answering systems to agents that can interpret intent, call tools, and report completed actions through ordinary messaging applications."
    )
    p = doc.add_paragraph()
    p.add_run(
        "The purpose of this project was not to reproduce or install OpenClaw. Instead, it used the concept as architectural inspiration and implemented a deliberately smaller system whose components, safety boundaries, and behavior could be understood and demonstrated by a single developer on consumer hardware."
    )
    doc.add_heading("1.1 Problem Statement", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "Academic and personal productivity work includes repetitive activities such as converting lecture PDFs into machine-readable text, reviewing email, checking deadlines, and combining tasks with calendar information. Existing automation platforms can solve parts of this problem, but they often require preconfigured workflows, broad platform adoption, or subscriptions. The project therefore explored whether a small natural-language agent could connect these tasks while retaining explicit safety controls."
    )
    doc.add_heading("1.2 Project Objectives", level=2)
    add_bullets(
        doc,
        [
            "Provide a Discord-based natural-language interface that is accessible only to the configured owner.",
            "Deliver an OCR-first MVP that creates LLM-readable text files from PDFs stored in OneDrive.",
            "Add practical Gmail, Google Sheets, and Google Calendar integrations.",
            "Demonstrate continuous operation on a MacBook Air M1 without Docker, a local LLM, MCP, browser automation, or a multi-agent framework.",
            "Keep test and demonstration API usage below KRW 10,000.",
        ],
    )

    doc.add_heading("2. Scope and Requirements", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "OCR was prioritized as the required vertical slice. Gmail management and the daily briefing were implemented after the OCR path was proven. The scope intentionally excluded persistent agent memory, arbitrary shell execution, hardware control, local-model deployment, and multi-user support."
    )
    add_requirements_table(doc)
    doc.add_heading("2.1 Explicitly Excluded Functions", level=2)
    add_bullets(
        doc,
        [
            "OpenClaw repository installation or reuse",
            "MCP and multi-agent architecture",
            "Browser automation and arbitrary terminal commands",
            "Calendar create, update, or delete operations",
            "Discord file transfer",
            "Long-term conversational memory or database-backed state",
        ],
    )

    doc.add_heading("3. System Architecture", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The architecture separates probabilistic language-model work from deterministic privileged actions. The model interprets commands, transcribes page images, classifies email, and drafts summaries. Python code—not the model—controls paths, Google API calls, output creation, confirmation state, and allowlisted tool selection."
    )
    pic = doc.add_picture(str(DIAGRAM), width=Inches(6.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    set_run(caption.add_run("Figure 1. Implemented system architecture and trust boundary."), 9, False, MUTED, True)

    doc.add_heading("3.1 Main Components", level=2)
    add_bullets(
        doc,
        [
            "Discord Bot: receives owner messages, rejects attachments, acknowledges processing, and splits long responses to fit Discord limits.",
            "Agent Controller: routes validated intents and maintains only short-lived OCR selection and Gmail confirmation state.",
            "OCR Tool: indexes PDF names inside the configured OneDrive root, renders selected pages, and sends images to the GPT vision-capable endpoint.",
            "Gmail Tool: performs date-bounded review, structured classification, conservative Trash actions, reply drafting, and explicit send/cancel confirmation.",
            "Sheets and Calendar Tools: read task records and upcoming events without write permissions.",
            "launchd Service: starts at login and restarts the bot if the process exits.",
        ],
    )

    doc.add_heading("4. Functional Implementation", level=1)
    doc.add_heading("4.1 Discord Natural-Language Interface", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "The bot uses both deterministic phrase matching and structured GPT parsing. Common commands are routed without an additional model call, while less predictable wording is mapped to a fixed intent schema. Only the configured 19-digit Discord user ID is accepted. Each accepted request receives an immediate “Command received. Processing...” acknowledgement."
    )

    doc.add_heading("4.2 PDF OCR Automation", level=2)
    add_numbered(
        doc,
        [
            "Search recursively within the configured OneDrive root using normalized fuzzy filename matching.",
            "Ask the owner to select a candidate when the file reference is ambiguous.",
            "Validate the requested page range and enforce the maximum pages per run.",
            "Render each PDF page to a JPEG image and transcribe it with the GPT API.",
            "Write Korean and English text, LaTeX mathematics, Markdown tables, figure descriptions, and page separators.",
            "Create `<original_name>_ag.txt` atomically beside the source PDF.",
        ],
    )
    add_callout(
        doc,
        "Overwrite guarantee",
        "The tool checks for an existing output before OCR and opens the final output in exclusive-create mode. If a matching `_ag.txt` already exists—or appears while processing—the operation stops without replacing it.",
    )

    doc.add_heading("4.3 Gmail Management", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "Gmail organization requires an explicit range such as “last 48 hours” or “2026-06-11 to 2026-06-12.” If the owner omits the range, the bot asks for it and performs no mailbox action. Email is classified as important, schedule-related, promotional, or other. Only unmistakable promotional messages may be moved to Gmail Trash."
    )
    add_bullets(
        doc,
        [
            "Protected categories include personal, work, school, account, payment, receipt, authentication, security, legal, deadline-related, and ambiguous email.",
            "Injection-like text inside an email prevents automatic Trash movement.",
            "The Discord response reports the reviewed count, Trash count, selected subjects, and important summaries.",
            "A Markdown report named `important_emails.md` is saved locally for evidence and review.",
        ],
    )

    doc.add_heading("4.4 Reply Draft and Approval", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "The agent can search for an email within an optional date range and generate a concise English reply draft. The draft is not sent immediately. It is assigned a short ID and displayed in Discord. The owner must enter `send <ID>` to transmit it or `cancel <ID>` to discard it. Pending drafts are held only in memory and disappear after restart."
    )

    doc.add_heading("4.5 Daily Productivity Briefing", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "The daily briefing combines important or schedule-related Gmail messages from the last 48 hours, unfinished Google Sheets tasks, and primary Google Calendar events from the next 48 hours. Completed rows and formula-only artifacts such as D+9699 are excluded. The final briefing is delivered directly in Discord rather than saved or attached as a file."
    )

    doc.add_heading("5. Security and Safety Design", level=1)
    add_callout(
        doc,
        "Security principle",
        "The LLM may interpret data, but it cannot directly execute shell commands, select arbitrary tools, escape configured paths, or send an email without the deterministic controller’s permission.",
        "FFF3CD",
    )
    add_bullets(
        doc,
        [
            "Owner-only Discord authorization using a configured user ID.",
            "Filesystem access restricted to one resolved OneDrive root; path traversal and symlink escape are rejected.",
            "PDF-only OCR input and Discord attachment rejection.",
            "Fixed Pydantic intent and decision schemas rather than free-form tool instructions.",
            "Untrusted PDF, email, Sheet, and Calendar content wrapped with explicit data-only delimiters.",
            "No use of `eval`, `exec`, model-generated shell commands, or arbitrary browser control.",
            "OpenAI API requests use `store=False`.",
            "Secrets remain in `.env`, `credentials.json`, and `token.json`, which are excluded from submission and restricted to owner-only filesystem permissions.",
        ],
    )

    doc.add_heading("6. Deployment and Continuous Operation", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The application runs in a Python virtual environment on macOS. Google services use OAuth desktop-app credentials. A launchd LaunchAgent includes `RunAtLoad` and `KeepAlive`, allowing the Discord bot to start automatically at login and restart after failure. AC-power system sleep was disabled so the Mac can remain available continuously while connected to power and the internet."
    )
    add_bullets(
        doc,
        [
            "Hardware: MacBook Air M1, 8 GB RAM, 256 GB SSD",
            "Runtime: Python, discord.py, OpenAI Python SDK, PyMuPDF, Google APIs",
            "Storage: OneDrive-mounted local filesystem plus local logs/reports",
            "Operational status: AgenticAI_Bot connected to Discord and managed by launchd",
        ],
    )

    doc.add_heading("7. Testing and Results", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "Testing combined automated unit tests with live end-to-end verification. The automated suite covers owner restriction, path safety, atomic output creation, PDF indexing, unreadable OneDrive placeholders, command routing, Gmail protection rules, date parsing, report writes, task completion filtering, and controller behavior."
    )
    add_test_table(doc)
    doc.add_heading("7.1 OCR Output Quality", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "The engineering mathematics assignment test preserved Korean instructions and mathematical expressions such as complex numbers, roots, powers, and exponential notation in LaTeX form. The timetable test verified general structured text extraction. Both outputs used the required `_ag.txt` suffix and contained explicit page markers."
    )
    doc.add_heading("7.2 Gmail Live Result", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "For the inclusive range June 11–12, 2026, the agent reviewed 10 inbox messages, moved 7 unmistakable promotional messages to Trash, retained important content, summarized a Coursera course message, and reported the results in Discord. A subsequent date-bounded reply-draft command successfully produced a draft and supported cancellation."
    )
    doc.add_heading("7.3 Observed Corrections During Integration", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "Live Discord testing exposed configuration and routing problems that unit-only testing would not have found, including a mismatched bot application, a server/channel setup error, and an undefined variable in the Gmail date-range path. Each problem was corrected, converted into a regression test where applicable, and redeployed through the persistent service."
    )

    doc.add_heading("8. Comparison with Existing Services", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The comparison below uses the project’s March 2026 frame and focuses on architectural fit rather than declaring a universal winner. Commercial products offer much broader integration catalogs and organizational governance. OpenClaw provides a more general personal-agent runtime. This project intentionally trades breadth for inspectability, a small attack surface, and exact control over the demonstrated workflows."
    )
    add_comparison_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    set_run(
        p.add_run(
            "Interpretation: the implemented MVP is not a replacement for these platforms. Its contribution is a transparent, narrow implementation that demonstrates the same core agentic loop—understand, select a tool, act, and report—while making high-risk decisions explicit."
        ),
        9.5,
        False,
        MUTED,
        True,
    )

    doc.add_heading("9. Limitations and Future Work", level=1)
    add_bullets(
        doc,
        [
            "The system depends on cloud APIs and internet connectivity.",
            "OneDrive online-only placeholders must be downloaded before OCR.",
            "OCR is processed page by page, so long documents increase cost and latency.",
            "Email classification remains probabilistic despite conservative deterministic checks.",
            "Pending confirmations are stored only in memory and are lost after process restart.",
            "The Mac must remain powered, logged in, connected, and awake.",
            "Python 3.9 currently produces dependency support warnings; migration to Python 3.11+ is recommended.",
        ],
    )
    doc.add_heading("9.1 Recommended Extensions", level=2)
    add_bullets(
        doc,
        [
            "Persistent, encrypted approval state with expiration times",
            "Per-action audit log and cost/latency dashboard",
            "Batch OCR with resumable checkpoints and quality scoring",
            "Calendar creation behind explicit confirmation",
            "Automated health checks and failure notifications",
            "Evaluation dataset for OCR accuracy and email false-positive rates",
        ],
    )

    doc.add_heading("10. Conclusion", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "The project achieved its objective of demonstrating a practical Agentic AI system on consumer hardware. A Discord message can now trigger real work across local files and Google services, while deterministic code controls permissions and sensitive actions. The OCR-first strategy produced an early demonstrable success, and the same architecture was extended to Gmail organization, approval-based replies, task records, calendar inspection, and daily briefings."
    )
    p = doc.add_paragraph()
    p.add_run(
        "The final system shows that a useful personal automation agent does not require a large framework or local model deployment. A small controller, a hosted reasoning model, narrowly scoped tools, explicit user approval, and careful filesystem and content boundaries were sufficient to produce a functioning MVP. The most important outcome is therefore both functional and architectural: meaningful automation can be achieved quickly without surrendering all control to an unrestricted autonomous runtime."
    )

    doc.add_heading("References", level=1)
    add_source(
        doc,
        1,
        "OpenClaw — Personal AI Assistant",
        "https://openclaw.ai/",
        "Official product page describing chat-based inbox, email, calendar, and personal-assistant automation.",
    )
    add_source(
        doc,
        2,
        "Zapier Agents",
        "https://zapier.com/agents",
        "Official page describing specialized agents and integrations across more than 9,000 applications.",
    )
    add_source(
        doc,
        3,
        "n8n AI Workflow Automation",
        "https://n8n.io/ai/",
        "Official page emphasizing business rules, auditability, and human-in-the-loop checkpoints.",
    )
    add_source(
        doc,
        4,
        "Microsoft 365 Copilot",
        "https://www.microsoft.com/en-us/microsoft-365-copilot",
        "Official page describing work-context chat, agents, Microsoft 365 integration, and governed access.",
    )
    add_source(
        doc,
        5,
        "Project source code and implementation specification",
        "Local project workspace, June 20–21, 2026",
        "Primary evidence for implemented behavior, tests, deployment, and live results.",
    )

    doc.add_heading("Appendix A. Demonstration Commands", level=1)
    commands = [
        "help",
        "OCR pages 1-3 of modern control systems",
        "Organize my emails",
        "last 48 hours",
        "Organize my emails from 2026-06-11 to 2026-06-12",
        "Draft a reply to the email from Coursera from 2026-06-11 to 2026-06-12",
        "cancel <ID>",
        "Show my calendar",
        "Create my daily briefing",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(command)
        set_run(run, 9.5, False, INK, font="Courier New")
        set_cell = None

    doc.add_heading("Appendix B. Submission Security Checklist", level=1)
    add_bullets(
        doc,
        [
            "Include source code, README, implementation specification, requirements, `.env.example`, proposal, final report, and demonstration video.",
            "Exclude `.env`, `credentials.json`, `token.json`, `.venv`, runtime logs, and personal OCR source files.",
            "Reset any credential that has been exposed outside the local machine.",
            "Verify the Discord bot is either intentionally left running or stopped after the presentation.",
        ],
    )

    # Widow/orphan and table row behavior.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    doc.core_properties.title = "Final Report - Agentic AI-Based Task Automation System"
    doc.core_properties.subject = "Implementation and evaluation of a personal Agentic AI automation MVP"
    doc.core_properties.author = "Project Developer"
    doc.core_properties.keywords = "Agentic AI, Discord, OCR, Gmail, Google Sheets, Google Calendar, OpenClaw"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
